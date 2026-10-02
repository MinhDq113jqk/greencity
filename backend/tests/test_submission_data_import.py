"""FCS-17 synthetic-pack acceptance on the disposable PostgreSQL runner."""

from __future__ import annotations

from datetime import UTC, datetime
from hashlib import sha256
import json
import os
from pathlib import Path
import secrets
import shutil

from openpyxl import load_workbook
import pytest
from sqlalchemy import func, select

from app.core.config import Settings
from app.core.database import Database
from app.models.billing import BillingInvoice, Payment
from app.models.building import Building
from app.models.site import Site
from app.models.submission_import import SubmissionExternalReference
from app.models.tenant import Tenant
from app.models.unit import Unit
from app.services.submission_data_import import SubmissionImportError
from app.services.submission_data_reader import read_submission_pack
from scripts.generate_submission_test_pack import generate_pack
from scripts.load_submission_data import apply, dry_run, verify
from scripts.submission_data_preflight import run_preflight


FIXTURE_ROOT = Path(__file__).parent / "fixtures" / "submission_data"
PACK = FIXTURE_ROOT / "synthetic"
MANIFEST = FIXTURE_ROOT / "synthetic_manifest.json"


def test_synthetic_pack_contract_privacy_and_finance_oracle():
    """The submitted fixture contains only generated surrogates and reconciles."""

    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    assert manifest["classification"] == "SYNTHETIC_TEST_ONLY"
    assert manifest["workbook_count"] == 12
    assert manifest["total_rows"] == 101
    pack = read_submission_pack(PACK)
    assert len(pack.workbooks) == 12 and pack.total_rows == 101
    preflight = run_preflight(PACK, as_of_utc=datetime(2026, 9, 26, 16, 59, 59, tzinfo=UTC), manifest_path=MANIFEST)
    assert preflight.status == "PASS"
    assert preflight.orphan_count == preflight.future_count == 0
    assert preflight.as_report()["database_write"] is False

    rows = {
        workbook.file_name: [dict(zip(workbook.headers, values, strict=True)) for values in workbook.rows]
        for workbook in pack.workbooks
    }
    for record in rows["02_danh_sach_nhan_vien.xlsx"]:
        assert str(record["full_name"]).startswith("Synthetic Account ")
        assert str(record["email"]).endswith("@example.invalid")
        assert str(record["username"]).startswith("syn_")
    for record in rows["04_cu_dan.xlsx"]:
        assert str(record["full_name"]).startswith("Synthetic Resident ")
        assert str(record["phone_masked"]).startswith("***")
        assert str(record["email_masked"]).endswith("@example.invalid")
    for record in rows["12_buu_pham.xlsx"]:
        assert str(record["recipient_name_snapshot"]).startswith("Synthetic Recipient ")
        assert str(record["recipient_contact_masked"]).startswith("***")

    invoices = {row["invoice_number"]: row for row in rows["10_hoa_don.xlsx"]}
    payments = rows["11_thanh_toan.xlsx"]
    assert len(invoices) == 8 and len(payments) == 8
    assert {row["status"] for row in payments} == {"ALLOCATED", "PARTIALLY_ALLOCATED", "OVERPAID", "UNMATCHED"}
    assert all(row["total_vnd"] == row["amount_vnd"] for row in invoices.values())
    for invoice in invoices.values():
        allocated = sum(row["allocated_vnd"] or 0 for row in payments if row["matched_invoice_number"] == invoice["invoice_number"])
        assert invoice["outstanding_vnd"] == invoice["total_vnd"] - allocated


def test_generator_rejects_raw_target_and_creates_200_unit_scale_fixture(tmp_path):
    with pytest.raises(ValueError, match="raw or approved"):
        generate_pack(tmp_path / "excel-data")
    manifest = generate_pack(tmp_path / "fixture")
    assert manifest["total_rows"] == 101
    assert len(read_submission_pack(tmp_path / "fixture" / "synthetic").workbooks) == 12
    lines = (tmp_path / "fixture" / "synthetic_scale_200_units.csv").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 201
    assert len({line.split(",")[3] for line in lines[1:]}) == 200


def _counts() -> tuple[int, int, int, int]:
    database = Database(Settings())
    try:
        with database.get_session() as session:
            tenant = session.scalar(select(Tenant).where(Tenant.code == "SYN-TENANT"))
            if tenant is None:
                return (0, 0, 0, 0)
            return (
                session.scalar(select(func.count(Unit.id)).join(Building, Unit.building_id == Building.id).join(Site, Building.site_id == Site.id).where(Site.tenant_id == tenant.id)) or 0,
                session.scalar(select(func.count(BillingInvoice.id)).where(BillingInvoice.tenant_id == tenant.id)) or 0,
                session.scalar(select(func.count(Payment.id)).where(Payment.tenant_id == tenant.id)) or 0,
                session.scalar(select(func.count(SubmissionExternalReference.id)).where(SubmissionExternalReference.tenant_id == tenant.id)) or 0,
            )
    finally:
        database.close()


@pytest.mark.integration
@pytest.mark.skipif(
    os.getenv("GREENCITY_ISOLATED_SECURITY_TESTS") != "1",
    reason="Requires scripts.test_isolated disposable PostgreSQL cluster",
)
def test_synthetic_pack_all_stage_apply_verify_retry_and_rollback(tmp_path, monkeypatch):
    settings = Settings()
    assert settings.app_env == "test"
    assert settings.sqlalchemy_url().host == "127.0.0.1"
    assert settings.sqlalchemy_url().username == "test_migrator"
    assert os.getenv("GREENCITY_DISABLE_DOTENV") == "1"

    pack = read_submission_pack(PACK)
    staff = next(workbook for workbook in pack.workbooks if workbook.file_name == "02_danh_sach_nhan_vien.xlsx")
    usernames = [dict(zip(staff.headers, values, strict=True))["username"] for values in staff.rows]
    credential_file = tmp_path / "synthetic-credentials.json"
    credential_file.write_text(
        json.dumps({username: secrets.token_urlsafe(32) for username in usernames}),
        encoding="utf-8",
    )
    monkeypatch.setenv("GREENCITY_SUBMISSION_CREDENTIALS_FILE", str(credential_file))
    monkeypatch.setenv("GREENCITY_SUBMISSION_PIN_SECRET", secrets.token_urlsafe(32))
    monkeypatch.setenv("GREENCITY_SUBMISSION_EVIDENCE_DIR", str(FIXTURE_ROOT / "synthetic_evidence"))
    monkeypatch.setenv("PRIVATE_STORAGE_PATH", str(tmp_path / "private-evidence"))

    assert _counts() == (0, 0, 0, 0)
    preview = dry_run(PACK, stage="all")
    assert preview["status"] == "PASS" and preview["database_write"] is False
    assert _counts() == (0, 0, 0, 0)

    invalid = tmp_path / "invalid"
    shutil.copytree(PACK, invalid)
    source = invalid / "01_cong_viec.xlsx"
    workbook = load_workbook(source)
    sheet = workbook["Dữ liệu"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("sla_due_at") + 1).value = sheet.cell(2, headers.index("created_at") + 1).value
    workbook.save(source)
    workbook.close()
    assert run_preflight(invalid, as_of_utc=datetime.now(UTC)).status == "PASS"
    invalid_manifest = tmp_path / "invalid_manifest.json"
    metadata = json.loads(MANIFEST.read_text(encoding="utf-8"))
    for entry in metadata["workbooks"]:
        if entry["file_name"] == source.name:
            entry["sha256"] = sha256(source.read_bytes()).hexdigest()
    invalid_manifest.write_text(json.dumps(metadata), encoding="utf-8")
    with pytest.raises(SubmissionImportError) as error:
        apply(invalid, stage="all", manifest_path=invalid_manifest)
    assert error.value.code == "CATEGORY_SLA_INVALID"
    assert _counts() == (0, 0, 0, 0)

    created = apply(PACK, stage="all")
    assert created["status"] == "PASS" and created["database_write"] is True
    assert created["counts"]["fcs16_payments"]["ar_delta_vnd"] == 0
    checked = verify(PACK, stage="all")
    assert checked["status"] == "PASS" and checked["reconciliation"]["master_units"] == 12
    assert checked["reconciliation"]["replay_invoices"] == 8
    assert checked["reconciliation"]["replay_payments"] == 8
    before_retry = _counts()
    assert before_retry[:3] == (12, 8, 8)
    repeated = apply(PACK, stage="all")
    assert repeated["status"] == "REPLAY" and repeated["database_write"] is False
    assert _counts() == before_retry
