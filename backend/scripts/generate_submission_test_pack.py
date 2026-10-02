"""Generate a wholly synthetic FCS-17 XLSX pack for local import rehearsal.

No source workbook is opened. The output is test data, not an approved
transformation of the user's private ``excel-data`` directory.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime, timedelta
from hashlib import sha256
import json
from pathlib import Path
import struct
import zlib

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.submission_data_contract import CONTRACT_VERSION, WORKBOOK_CONTRACTS


PACK_VERSION = "synthetic-fcs17-v1"
TENANT = "SYN-TENANT"
SITE = "SYN-SITE"
BUILDING = "SYN-B1"
PERIOD = "2026-07"
RATE = 10_000
ROUNDING = 1_000


def _at(day: int, hour: int = 9, minute: int = 0, *, month: int = 7) -> datetime:
    return datetime(2026, month, day, hour, minute)


def _unit(index: int) -> str:
    return f"SYN-U{index:03d}"


def _account(index: int) -> str:
    return f"SYN-BA-{index:03d}"


def _invoice(index: int) -> str:
    return f"SYN-INV-{index:03d}"


def _scope() -> dict[str, str]:
    return {"tenant_code": TENANT, "site_code": SITE, "building_code": BUILDING}


def _build_rows() -> dict[str, list[dict[str, object]]]:
    result: dict[str, list[dict[str, object]]] = {}

    roles = [
        "admin", "director", "cskh", "accountant", "technical_lead",
        "technician", "technician", "cleaning", "cleaning", "security",
        "security", "resident", "resident", "resident", "resident",
        "resident", "resident", "resident",
    ]
    role_number: dict[str, int] = {}
    users: dict[str, list[str]] = {}
    staff: list[dict[str, object]] = []
    for index, role in enumerate(roles, start=1):
        role_number[role] = role_number.get(role, 0) + 1
        username = f"syn_{role}_{role_number[role]:02d}"
        users.setdefault(role, []).append(username)
        staff.append({
            "employee_code": f"SYN-EMP-{index:03d}",
            "full_name": f"Synthetic Account {index:02d}",
            "username": username,
            "email": f"synthetic{index:02d}@example.invalid",
            "phone": f"***{index:04d}",
            "role": role,
            "department": "Synthetic Test",
            "tenant_code": TENANT,
            "site_code": SITE,
            "building_code": None if role == "director" else BUILDING,
            "status": "ACTIVE",
            "start_date": date(2026, 1, 1),
        })
    result["02_danh_sach_nhan_vien.xlsx"] = staff

    units: list[dict[str, object]] = []
    for index in range(1, 13):
        units.append({
            **_scope(), "site_name": "Synthetic Site", "building_name": "Synthetic Building",
            "floors_count": 12, "unit_number": _unit(index), "floor": index,
            "area_m2": 45 + 5 * index,
            "unit_status": "occupied" if index <= 10 else "vacant",
            "billing_account_number": _account(index),
            "billing_status": "ACTIVE" if index <= 8 else "SUSPENDED",
            "resident_code": f"SYN-RES-{index:03d}" if index <= 10 else None,
        })
    result["03_toa_nha_can_ho.xlsx"] = units

    residents: list[dict[str, object]] = []
    for index in range(1, 12):
        owner = index <= 10
        residents.append({
            "resident_code": f"SYN-RES-{index:03d}",
            "full_name": f"Synthetic Resident {index:02d}",
            "phone_masked": f"***{index:04d}",
            "email_masked": f"r***{index:02d}@example.invalid",
            "relationship_type": "owner" if owner else "family_member",
            "ownership_ratio": 1 if owner else 0,
            **_scope(), "unit_number": _unit(index if owner else 1),
            "valid_from": date(2026, 1, 1), "valid_to": None,
            "status": "ACTIVE",
        })
    result["04_cu_dan.xlsx"] = residents

    works: list[dict[str, object]] = []
    work_statuses = ["NEW", "IN_PROGRESS", "CLOSED", "IN_PROGRESS", "NEW", "IN_PROGRESS", "NEW", "IN_PROGRESS", "IN_PROGRESS"]
    for index, status in enumerate(work_statuses, start=1):
        created = _at(10 + index)
        works.append({
            "work_code": f"SYN-WORK-{index:03d}",
            "request_type": "SYN-MAINT",
            "title": f"Synthetic work {index:02d}",
            "description": "Synthetic maintenance request for local rehearsal.",
            "priority": ("LOW", "MEDIUM", "HIGH")[index % 3],
            "status": status,
            **_scope(), "unit_number": _unit(index),
            "reported_by_username": users["cskh"][0],
            "assignee_username": users["technician"][(index - 1) % 2],
            "created_at": created,
            "sla_due_at": created + timedelta(hours=4),
            "completed_at": created + timedelta(hours=2) if status == "CLOSED" else None,
            "cost_vnd": 25_000 * index,
        })
    result["01_cong_viec.xlsx"] = works

    assets: list[dict[str, object]] = []
    for index in range(1, 7):
        assets.append({
            "asset_code": f"SYN-ASSET-{index:03d}",
            "asset_name": f"Synthetic asset {index:02d}",
            "asset_type": "Synthetic equipment",
            **_scope(), "location": f"Synthetic floor {index:02d}",
            "status": "ACTIVE", "installed_on": date(2025, 1, 1),
            "maintenance_plan_code": f"SYN-PLAN-{index:03d}",
            "interval_days": 30, "next_due_date": date(2026, 7, 1 + index),
            "vendor_name": "Synthetic Vendor",
            "responsible_username": users["technician"][(index - 1) % 2],
        })
    result["05_tai_san_bao_tri.xlsx"] = assets

    cleaning: list[dict[str, object]] = []
    cleaning_statuses = ["PLANNED", "IN_PROGRESS", "COMPLETED", "COMPLETED", "COMPLETED", "CANCELLED"]
    for index, status in enumerate(cleaning_statuses, start=1):
        failed = index == 5
        checked = "FAIL" if failed else "PASS" if status == "COMPLETED" else "PENDING"
        cleaning.append({
            "shift_code": f"SYN-CLEAN-{index:03d}",
            "route_code": f"SYN-ROUTE-{index:03d}",
            "area_code": f"SYN-AREA-{index:03d}",
            "site_code": SITE, "building_code": BUILDING,
            "scheduled_start": _at(index, 8), "scheduled_end": _at(index, 10),
            "assignee_username": users["cleaning"][(index - 1) % 2],
            "status": status,
            "checklist_floor": checked,
            "checklist_bins": checked,
            "quality_result": checked,
            "rework_required": failed,
        })
    result["06_ve_sinh.xlsx"] = cleaning

    incidents: list[dict[str, object]] = []
    incident_statuses = ["NEW", "TRIAGED", "IN_PROGRESS", "RESOLVED", "RESOLVED", "NEW"]
    for index, status in enumerate(incident_statuses, start=1):
        incidents.append({
            "incident_code": f"SYN-INC-{index:03d}",
            "incident_type": "FIRE" if index == 5 else "SECURITY",
            "severity": ("LOW", "MEDIUM", "HIGH", "LOW", "MEDIUM", "LOW")[index - 1],
            "status": status, **_scope(),
            "location": f"Synthetic location {index:02d}",
            "reported_at": _at(index, 11),
            "reported_by_username": users["security"][(index - 1) % 2],
            "owner_username": users["security"][(index - 1) % 2],
            "resolved_at": _at(index, 12) if status == "RESOLVED" else None,
            "description": "Synthetic security incident for local rehearsal.",
        })
    result["08_su_co_an_ninh.xlsx"] = incidents

    patrols: list[dict[str, object]] = []
    patrol_statuses = ["COMPLETED", "COMPLETED", "SCHEDULED", "CANCELLED", "SCHEDULED", "CANCELLED"]
    for index, status in enumerate(patrol_statuses, start=1):
        patrols.append({
            "shift_code": f"SYN-PATROL-{index:03d}",
            "site_code": SITE, "building_code": BUILDING,
            "guard_username": users["security"][(index - 1) % 2],
            "scheduled_start": _at(index, 18),
            "scheduled_end": _at(index, 20),
            "patrol_point_code": f"SYN-POINT-{index:03d}",
            "window_start": _at(index, 18, 30),
            "window_end": _at(index, 19),
            "patrol_status": status,
            "event_type": "CHECK_IN" if status == "COMPLETED" else None,
            "incident_code": "SYN-INC-001" if index == 1 else None,
            "handoff_summary": "Synthetic shift handoff." if index == 1 else None,
        })
    result["07_an_ninh_tuan_tra.xlsx"] = patrols

    policy_codes = ["FEE-MONTHLY", "AUX-WATER", "AUX-WASTE", "AUX-PARK", "AUX-OTHER"]
    months = [7, 8, 4, 6, 5]
    policies: list[dict[str, object]] = []
    for code, month in zip(policy_codes, months, strict=True):
        next_month = date(2026, month + 1, 1)
        period_end = next_month - timedelta(days=1)
        policies.append({
            "fee_policy_code": code,
            "fee_policy_name": f"Synthetic {code.lower()} policy",
            **_scope(), "version_number": 1,
            "effective_from": date(2026, month, 1), "effective_to": None,
            "unit_rate_vnd": RATE if code == "FEE-MONTHLY" else 1_000,
            "basis": "UNIT_AREA_M2", "rounding_unit_vnd": ROUNDING,
            "period_key": f"2026-{month:02d}",
            "period_start": date(2026, month, 1), "period_end": period_end,
            "cutoff_at": datetime(2026, month, period_end.day, 23, 0),
            "period_status": "OPEN", "policy_status": "ACTIVE",
        })
    result["09_chinh_sach_phi.xlsx"] = policies

    invoices: list[dict[str, object]] = []
    payment_amounts = [500_000, 550_000, 600_000, 300_000, 700_000, 750_000, 850_000]
    for index in range(1, 9):
        area = 45 + 5 * index
        total = area * RATE
        paid = min(payment_amounts[index - 1], total) if index <= 7 else 0
        invoice_status = "ISSUED" if paid == 0 else "PAID" if paid == total else "PARTIALLY_PAID"
        invoices.append({
            "invoice_number": _invoice(index),
            "billing_account_number": _account(index),
            **_scope(), "unit_number": _unit(index),
            "period_key": PERIOD, "billing_run_key": "SYN-RUN-2026-07",
            "policy_version": 1, "invoice_status": invoice_status,
            "issued_on": date(2026, 8, 1), "due_on": date(2026, 8, 15),
            "total_vnd": total, "outstanding_vnd": total - paid,
            "line_code": "FEE", "line_description": "UNIT AREA M2",
            "basis_quantity": area, "unit_rate_vnd_snapshot": RATE,
            "rounding_unit_vnd_snapshot": ROUNDING, "amount_vnd": total,
        })
    result["10_hoa_don.xlsx"] = invoices

    payments: list[dict[str, object]] = []
    for index, amount in enumerate(payment_amounts, start=1):
        total = (45 + 5 * index) * RATE
        payments.append({
            "source_reference": f"SYN-PAY-{index:03d}",
            "receipt_number": f"SYN-REC-{index:03d}",
            "payment_source": ("CASH", "BANK_TRANSFER", "GATEWAY")[(index - 1) % 3],
            **_scope(), "billing_account_number": _account(index),
            "received_at": _at(10, 10 + index, month=8),
            "amount_vnd": amount,
            "status": "OVERPAID" if amount > total else "PARTIALLY_ALLOCATED" if amount < total else "ALLOCATED",
            "matched_invoice_number": _invoice(index),
            "allocated_vnd": min(amount, total),
            "unmatched_reason": None,
            "overpayment_vnd": max(amount - total, 0),
            "received_by_username": users["accountant"][0],
        })
    payments.append({
        "source_reference": "SYN-PAY-008", "receipt_number": "SYN-REC-008",
        "payment_source": "BANK_TRANSFER", **_scope(),
        "billing_account_number": None,
        "received_at": _at(10, 18, month=8), "amount_vnd": 45_000,
        "status": "UNMATCHED", "matched_invoice_number": None,
        "allocated_vnd": None, "unmatched_reason": "Synthetic unmatched receipt",
        "overpayment_vnd": None,
        "received_by_username": users["accountant"][0],
    })
    result["11_thanh_toan.xlsx"] = payments

    parcel_statuses = ["RECEIVED", "READY_FOR_PICKUP", "HANDED_OVER", "RETURNED", "LOST", "DAMAGED"]
    parcels: list[dict[str, object]] = []
    for index, status in enumerate(parcel_statuses, start=1):
        parcels.append({
            "parcel_code": f"SYN-PARCEL-{index:03d}",
            "site_code": SITE, "building_code": BUILDING,
            "unit_number": _unit(index),
            "recipient_name_snapshot": f"Synthetic Recipient {index:02d}",
            "recipient_contact_masked": f"***{index:04d}",
            "received_at": _at(index, 8, month=8),
            "status": status, "storage_location": f"SYN-SHELF-{index:02d}",
            "ready_at": _at(index, 9, month=8) if status != "RECEIVED" else None,
            "handed_over_at": _at(index, 10, month=8) if status == "HANDED_OVER" else None,
            "pin_attempt_count": 0,
            "case_required": status in {"RETURNED", "LOST", "DAMAGED"},
        })
    result["12_buu_pham.xlsx"] = parcels
    return result


def _write_workbook(path: Path, rows: list[dict[str, object]], contract) -> None:
    workbook = Workbook()
    data = workbook.active
    data.title = contract.data_sheet
    data.append([field.source_name for field in contract.fields])
    for row in rows:
        data.append([row.get(field.source_name) for field in contract.fields])
    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions
    guide = workbook.create_sheet(contract.guide_sheet)
    guide.append(["Pack", "Class", "Version", "Purpose"])
    guide.append(["GreenCity", "Synthetic test only", PACK_VERSION, "Local import rehearsal"])
    header_fill = PatternFill("solid", fgColor="DCEBDD")
    for sheet in workbook:
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="17351B")
            cell.fill = header_fill
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
    workbook.properties.creator = "GreenCity synthetic test generator"
    workbook.properties.title = "Synthetic GreenCity submission test data"
    workbook.properties.subject = "Synthetic test only; no real source cells"
    workbook.save(path)
    workbook.close()


def _synthetic_png() -> bytes:
    """An original green 8x8 PNG, with no personal or source-derived content."""

    def chunk(kind: bytes, payload: bytes) -> bytes:
        body = kind + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    header = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0)
    pixel_rows = (b"\x00" + b"\x45\xa0\x62" * 8) * 8
    return header + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(pixel_rows, 9)) + chunk(b"IEND", b"")


def generate_pack(output_root: Path) -> dict[str, object]:
    output_root = output_root.resolve()
    forbidden = {"excel-data", "data_that", "approved_pack"}
    if any(part.casefold() in forbidden for part in output_root.parts):
        raise ValueError("Refusing to write into a raw or approved data directory")
    pack_dir = output_root / "synthetic"
    pack_dir.mkdir(parents=True, exist_ok=True)
    existing = tuple(pack_dir.iterdir())
    if existing:
        manifest_path = output_root / "synthetic_manifest.json"
        if not manifest_path.is_file():
            raise FileExistsError("Existing XLSX files have no synthetic manifest")
        previous = json.loads(manifest_path.read_text(encoding="utf-8"))
        expected_names = {contract.file_name for contract in WORKBOOK_CONTRACTS}
        if previous.get("classification") != "SYNTHETIC_TEST_ONLY" or {path.name for path in existing} != expected_names:
            raise FileExistsError("Refusing to overwrite an unrecognized pack")
        previous_entries = {entry["file_name"]: entry["sha256"] for entry in previous["workbooks"]}
        if any(sha256(path.read_bytes()).hexdigest() != previous_entries.get(path.name) for path in existing):
            raise FileExistsError("Refusing to overwrite a modified synthetic workbook")
    rows_by_name = _build_rows()
    entries: list[dict[str, object]] = []
    for contract in WORKBOOK_CONTRACTS:
        rows = rows_by_name[contract.file_name]
        if len(rows) != contract.expected_data_rows:
            raise ValueError(f"Synthetic row count mismatch: {contract.file_name}")
        path = pack_dir / contract.file_name
        _write_workbook(path, rows, contract)
        entries.append({
            "file_name": contract.file_name,
            "sha256": sha256(path.read_bytes()).hexdigest(),
            "row_count": len(rows),
        })
    manifest = {
        "classification": "SYNTHETIC_TEST_ONLY",
        "pack_version": PACK_VERSION,
        "schema_version": CONTRACT_VERSION,
        "workbook_count": len(entries),
        "total_rows": sum(entry["row_count"] for entry in entries),
        "workbooks": entries,
    }
    (output_root / "synthetic_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    evidence_dir = output_root / "synthetic_evidence"
    evidence_dir.mkdir(exist_ok=True)
    evidence_path = evidence_dir / "SYN-WORK-003.png"
    if evidence_path.exists() and evidence_path.read_bytes() != _synthetic_png():
        raise FileExistsError("Refusing to overwrite modified synthetic evidence")
    evidence_path.write_bytes(_synthetic_png())
    scale_path = output_root / "synthetic_scale_200_units.csv"
    if scale_path.exists():
        with scale_path.open(newline="", encoding="utf-8") as handle:
            scale_rows = list(csv.reader(handle))
        if len(scale_rows) != 201 or scale_rows[0] != ["tenant_code", "site_code", "building_code", "unit_number", "floor", "area_m2"]:
            raise FileExistsError("Refusing to overwrite modified synthetic scale fixture")
    with scale_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(("tenant_code", "site_code", "building_code", "unit_number", "floor", "area_m2"))
        for index in range(1, 201):
            writer.writerow((TENANT, SITE, BUILDING, _unit(index), (index - 1) // 10 + 1, 40 + index % 25))
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate wholly synthetic GreenCity FCS-17 test data")
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()
    manifest = generate_pack(args.output_root)
    print(f"FCS17_SYNTHETIC=PASS workbooks={manifest['workbook_count']} rows={manifest['total_rows']} scale_units=200 evidence_images=1")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
