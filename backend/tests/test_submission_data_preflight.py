"""Read-only contract, mapping and preflight checks for the submission pack."""

from datetime import UTC, datetime
from hashlib import sha256
import json
from pathlib import Path
import shutil

import pytest

from app.services.submission_data_mapping import (
    CanonicalDerivationError,
    derive_fee_policy_code,
    derive_period_key,
    open_decision_ids,
    validate_mapping_catalog,
)
from app.services.submission_data_reader import (
    ReadLimits,
    SubmissionReaderError,
    read_submission_pack,
)
from app.services import submission_data_reader as reader
from scripts.submission_data_preflight import run_preflight


SOURCE = Path(__file__).resolve().parent / "fixtures" / "submission_data" / "synthetic"
SOURCE_MANIFEST = SOURCE.parent / "synthetic_manifest.json"
AS_OF = datetime(2026, 9, 26, 16, 59, 59, tzinfo=UTC)


def _rows(name: str) -> list[dict[str, object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(SOURCE / name, read_only=True, data_only=True)
    try:
        values = list(workbook["Dữ liệu"].values)
    finally:
        workbook.close()
    return [dict(zip(values[0], row, strict=True)) for row in values[1:]]


def test_fcs05_mapping_catalog_is_resolved_without_open_decisions():
    assert validate_mapping_catalog() == ()
    assert open_decision_ids() == ()


def test_canonical_invoice_and_unmatched_period_keys_are_deterministic():
    policies = _rows("09_chinh_sach_phi.xlsx")
    invoices = _rows("10_hoa_don.xlsx")
    payments = _rows("11_thanh_toan.xlsx")
    assert derive_fee_policy_code(invoices[0], policies) == "FEE-MONTHLY"
    assert derive_period_key(payments[1], invoices, policies) == "2026-07"
    assert derive_period_key(payments[-1], invoices, policies) == "2026-08"
    with pytest.raises(CanonicalDerivationError):
        derive_fee_policy_code({**invoices[0], "line_code": "UNKNOWN"}, policies)


def test_reader_accepts_the_12_workbook_envelope_without_db_access():
    pack = read_submission_pack(SOURCE)
    assert len(pack.workbooks) == 12
    assert pack.total_rows == 101
    assert all(workbook.source_size_bytes > 0 for workbook in pack.workbooks)


def test_reader_hash_covers_the_exact_rows_parsed_when_source_changes(tmp_path, monkeypatch):
    """A source edit after snapshotting cannot replace manifest-checked rows."""

    source = tmp_path / "01_cong_viec.xlsx"
    shutil.copyfile(SOURCE / source.name, source)
    expected = reader.read_submission_workbook(source)
    title_column = expected.headers.index("title")
    original_loader = reader.load_workbook
    changed = False

    def change_source_before_parse(file, **kwargs):
        nonlocal changed
        changed = True
        workbook = original_loader(source)
        workbook["Dữ liệu"].cell(row=2, column=title_column + 1).value = "Changed synthetic title"
        workbook.save(source)
        workbook.close()
        return original_loader(file, **kwargs)

    monkeypatch.setattr(reader, "load_workbook", change_source_before_parse)
    observed = reader.read_submission_workbook(source)
    assert changed
    assert sha256(source.read_bytes()).hexdigest() != expected.source_sha256
    assert observed.source_sha256 == expected.source_sha256
    assert observed.rows == expected.rows


def test_reader_rejects_a_tighter_resource_limit_without_disclosing_cells():
    with pytest.raises(SubmissionReaderError) as error:
        read_submission_pack(SOURCE, limits=ReadLimits(max_workbook_bytes=100))
    assert error.value.code == "WORKBOOK_TOO_LARGE"
    assert "Đèn" not in str(error.value)


def test_reader_rejects_formula_hidden_sheet_and_header_tampering(tmp_path):
    """Structural failures stay stable and never expose workbook cell values."""

    from openpyxl import load_workbook

    def copy_pack(name: str) -> Path:
        destination = tmp_path / name
        shutil.copytree(SOURCE, destination)
        return destination

    formula_pack = copy_pack("formula")
    workbook = load_workbook(formula_pack / "01_cong_viec.xlsx")
    workbook["Dữ liệu"]["A2"] = "=1+1"
    workbook.save(formula_pack / "01_cong_viec.xlsx")
    workbook.close()
    with pytest.raises(SubmissionReaderError) as formula_error:
        read_submission_pack(formula_pack)
    assert formula_error.value.code == "FORMULA_NOT_ALLOWED"

    hidden_pack = copy_pack("hidden")
    workbook = load_workbook(hidden_pack / "01_cong_viec.xlsx")
    workbook["Hướng dẫn"].sheet_state = "hidden"
    workbook.save(hidden_pack / "01_cong_viec.xlsx")
    workbook.close()
    with pytest.raises(SubmissionReaderError) as hidden_error:
        read_submission_pack(hidden_pack)
    assert hidden_error.value.code == "HIDDEN_SHEET_NOT_ALLOWED"

    header_pack = copy_pack("header")
    workbook = load_workbook(header_pack / "01_cong_viec.xlsx")
    workbook["Dữ liệu"]["A1"] = "unexpected_header"
    workbook.save(header_pack / "01_cong_viec.xlsx")
    workbook.close()
    with pytest.raises(SubmissionReaderError) as header_error:
        read_submission_pack(header_pack)
    assert header_error.value.code == "HEADER_MISMATCH"


def test_preflight_is_read_only_and_reports_zero_cross_file_orphans():
    result = run_preflight(SOURCE, as_of_utc=AS_OF)
    assert result.total_rows == 101
    assert result.decided_rows == 101
    assert result.orphan_count == 0
    assert result.as_report()["database_write"] is False
    assert not any(code.startswith("OPEN_MAPPING") for code in result.issue_counts)
    assert not any(code.startswith("ORPHAN_") for code in result.issue_counts)
    assert result.future_terminal_candidates == 0
    assert result.status == "PASS"
    assert result.source_checksum_verified is False


def test_preflight_verifies_complete_synthetic_and_raw_style_manifests(tmp_path):
    synthetic = run_preflight(SOURCE, as_of_utc=AS_OF, manifest_path=SOURCE_MANIFEST)
    assert synthetic.status == "PASS" and synthetic.source_checksum_verified

    manifest = json.loads(SOURCE_MANIFEST.read_text(encoding="utf-8"))
    for entry in manifest["workbooks"]:
        row_count = entry.pop("row_count")
        entry["size_bytes"] = (SOURCE / entry["file_name"]).stat().st_size
        entry["sheets"] = [{"sheet_role": "data", "data_row_count": row_count}]
    path = tmp_path / "raw-style-metadata.json"
    path.write_text(json.dumps(manifest), encoding="utf-8")
    raw_style = run_preflight(SOURCE, as_of_utc=AS_OF, manifest_path=path)
    assert raw_style.status == "PASS" and raw_style.source_checksum_verified


@pytest.mark.parametrize(
    ("change", "expected_code"),
    [
        ("duplicate", "SOURCE_MANIFEST_INVALID"),
        ("extra", "SOURCE_MANIFEST_INVALID"),
        ("missing", "SOURCE_MANIFEST_INVALID"),
        ("wrong_name", "SOURCE_MANIFEST_INVALID"),
        ("wrong_total", "SOURCE_MANIFEST_INVALID"),
        ("wrong_summary", "SOURCE_MANIFEST_INVALID"),
        ("row_count", "SOURCE_CHECKSUM_MISMATCH"),
        ("conflicting_sheet_count", "SOURCE_CHECKSUM_MISMATCH"),
        ("checksum", "SOURCE_CHECKSUM_MISMATCH"),
        ("size", "SOURCE_CHECKSUM_MISMATCH"),
        ("null_size", "SOURCE_CHECKSUM_MISMATCH"),
        ("invalid_digest", "SOURCE_CHECKSUM_MISMATCH"),
    ],
)
def test_preflight_rejects_incomplete_or_changed_manifest(tmp_path, change, expected_code):
    manifest = json.loads(SOURCE_MANIFEST.read_text(encoding="utf-8"))
    entries = manifest["workbooks"]
    if change == "duplicate":
        entries[-1] = dict(entries[0])
    elif change == "extra":
        entries.append({"file_name": "extra.xlsx", "sha256": "0" * 64, "row_count": 1})
    elif change == "missing":
        entries.pop()
    elif change == "wrong_name":
        entries[0]["file_name"] = "unknown.xlsx"
    elif change == "wrong_total":
        manifest["total_rows"] += 1
    elif change == "wrong_summary":
        manifest["summary"] = {"workbook_count": 12, "data_row_count": 102}
    elif change == "row_count":
        entries[0]["row_count"] += 1
    elif change == "conflicting_sheet_count":
        entries[0]["sheets"] = [{"sheet_role": "data", "data_row_count": 100}]
    elif change == "checksum":
        entries[0]["sha256"] = "0" * 64
    elif change == "size":
        entries[0]["size_bytes"] = 1
    elif change == "null_size":
        entries[0]["size_bytes"] = None
    else:
        entries[0]["sha256"] = "z" * 64
    path = tmp_path / "invalid-manifest.json"
    path.write_text(json.dumps(manifest), encoding="utf-8")

    result = run_preflight(SOURCE, as_of_utc=AS_OF, manifest_path=path)
    assert result.status == "FAIL"
    assert result.source_checksum_verified is False
    assert result.issue_counts.get(expected_code) == 1


def test_future_schedule_is_warning_only_and_does_not_block_preflight(tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "schedule-warning"
    shutil.copytree(SOURCE, source)
    workbook_path = source / "06_ve_sinh.xlsx"
    workbook = load_workbook(workbook_path)
    sheet = workbook["Dữ liệu"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("scheduled_start") + 1).value = datetime(2027, 1, 1, 8, 0)
    sheet.cell(2, headers.index("scheduled_end") + 1).value = datetime(2027, 1, 1, 9, 0)
    workbook.save(workbook_path)
    workbook.close()

    result = run_preflight(source, as_of_utc=datetime(2026, 9, 26, 16, 59, 59, tzinfo=UTC))
    assert result.status == "PASS"
    assert result.issue_counts == {"FUTURE_SCHEDULE": 2}
    assert {issue.severity for issue in result.issues} == {"WARNING"}
    assert result.as_report()["database_write"] is False


@pytest.mark.parametrize(
    ("file_name", "column", "row_number"),
    [
        ("01_cong_viec.xlsx", "created_at", 2),
        ("08_su_co_an_ninh.xlsx", "reported_at", 2),
        ("11_thanh_toan.xlsx", "received_at", 2),
        ("12_buu_pham.xlsx", "received_at", 2),
        ("12_buu_pham.xlsx", "ready_at", 3),
    ],
)
def test_future_historical_event_blocks_preflight(tmp_path, file_name, column, row_number):
    from openpyxl import load_workbook

    source = tmp_path / "future-historical-event"
    shutil.copytree(SOURCE, source)
    workbook_path = source / file_name
    workbook = load_workbook(workbook_path)
    sheet = workbook["Dữ liệu"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(row_number, headers.index(column) + 1).value = datetime(2027, 1, 1, 8, 0)
    workbook.save(workbook_path)
    workbook.close()

    result = run_preflight(source, as_of_utc=datetime(2026, 9, 26, 16, 59, 59, tzinfo=UTC))
    assert result.status == "FAIL"
    assert result.issue_counts == {"FUTURE_HISTORICAL_EVENT": 1}
    assert result.future_count == 1
    assert result.future_terminal_candidates == 0
    assert [(issue.file_name, issue.row_number, issue.column, issue.severity) for issue in result.issues] == [
        (file_name, row_number, column, "ERROR"),
    ]


@pytest.mark.parametrize(
    ("file_name", "column", "row_number"),
    [
        ("01_cong_viec.xlsx", "completed_at", 4),
        ("08_su_co_an_ninh.xlsx", "resolved_at", 5),
        ("12_buu_pham.xlsx", "handed_over_at", 4),
    ],
)
def test_future_terminal_event_still_blocks_preflight(tmp_path, file_name, column, row_number):
    from openpyxl import load_workbook

    source = tmp_path / "future-terminal-event"
    shutil.copytree(SOURCE, source)
    workbook_path = source / file_name
    workbook = load_workbook(workbook_path)
    sheet = workbook["Dữ liệu"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(row_number, headers.index(column) + 1).value = datetime(2027, 1, 1, 8, 0)
    workbook.save(workbook_path)
    workbook.close()

    result = run_preflight(source, as_of_utc=datetime(2026, 9, 26, 16, 59, 59, tzinfo=UTC))
    assert result.status == "FAIL"
    assert result.issue_counts == {"FUTURE_TERMINAL_EVENT": 1}
    assert result.future_count == 1
    assert result.future_terminal_candidates == 1
