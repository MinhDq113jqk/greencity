"""The distribution-pack builder must stay fail closed and keep no raw cells."""

from __future__ import annotations

from datetime import UTC, datetime
from hashlib import sha256
import json
from pathlib import Path
import shutil

from openpyxl import load_workbook
import pytest

from app.services.submission_data_reader import read_submission_pack
from scripts.anonymize_submission_pack import AnonymizationError, Surrogates, anonymize_pack
from scripts.submission_data_preflight import run_preflight


ROOT = Path(__file__).parent / "fixtures" / "submission_data"
SOURCE = ROOT / "synthetic"
SOURCE_MANIFEST = ROOT / "synthetic_manifest.json"
AS_OF = datetime(2026, 9, 27, 16, 59, 59, tzinfo=UTC)


def _copy_source(tmp_path: Path) -> tuple[Path, Path]:
    source = tmp_path / "source"
    shutil.copytree(SOURCE, source)
    manifest = tmp_path / "source_manifest.json"
    shutil.copyfile(SOURCE_MANIFEST, manifest)
    return source, manifest


def _refresh_manifest_entry(manifest_path: Path, workbook_path: Path) -> None:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for entry in manifest["workbooks"]:
        if entry["file_name"] == workbook_path.name:
            entry["sha256"] = sha256(workbook_path.read_bytes()).hexdigest()
            break
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")


def test_anonymized_pack_preserves_reconciliation_shape_and_removes_identity(tmp_path):
    source, manifest_path = _copy_source(tmp_path)
    staff_path = source / "02_danh_sach_nhan_vien.xlsx"
    workbook = load_workbook(staff_path)
    data = workbook["Dữ liệu"]
    headers = [cell.value for cell in data[1]]
    data.cell(2, headers.index("full_name") + 1).value = "Private marker person"
    data.cell(2, headers.index("phone") + 1).value = "0901234567"
    workbook.save(staff_path)
    workbook.close()
    _refresh_manifest_entry(manifest_path, staff_path)

    output = tmp_path / "approved_pack"
    manifest = anonymize_pack(source, manifest_path, output, as_of_utc=AS_OF)
    assert manifest["classification"] == "LOCAL_VALIDATION_ONLY"
    assert manifest["distribution_allowed"] is False
    assert manifest["source_classification"] == "SYNTHETIC_TEST_ONLY"
    assert manifest["workbook_count"] == 12 and manifest["total_rows"] == 101
    assert manifest["source_preflight"]["warning_count"] == 0
    assert len(read_submission_pack(output).workbooks) == 12
    checked = run_preflight(output, as_of_utc=AS_OF,
                            manifest_path=tmp_path / "local_validation_manifest.json")
    assert checked.status == "PASS" and checked.source_checksum_verified
    assert checked.orphan_count == 0 and checked.total_rows == 101

    staff = load_workbook(output / "02_danh_sach_nhan_vien.xlsx", read_only=True, data_only=True)
    try:
        rows = list(staff["Dữ liệu"].values)
        row = dict(zip(rows[0], rows[1], strict=True))
        assert row["full_name"].startswith("Synthetic Account-")
        assert row["email"].endswith("@example.invalid")
        assert row["phone"].startswith("***")
        assert row["full_name"] != "Private marker person"
        assert row["phone"] != "0901234567"
    finally:
        staff.close()
    assert b"Private marker person" not in (output / "02_danh_sach_nhan_vien.xlsx").read_bytes()
    assert "0901234567" not in (tmp_path / "local_validation_manifest.json").read_text(encoding="utf-8")


def test_raw_derived_pack_is_local_validation_only(tmp_path):
    source, manifest_path = _copy_source(tmp_path)
    source_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_manifest["classification"] = "RESTRICTED_REAL_LOCAL_ONLY"
    manifest_path.write_text(json.dumps(source_manifest), encoding="utf-8")

    output = tmp_path / "local_validation_pack"
    result = anonymize_pack(source, manifest_path, output, as_of_utc=AS_OF)

    assert result["classification"] == "LOCAL_VALIDATION_ONLY"
    assert result["distribution_allowed"] is False
    assert "Original amounts" in result["privacy_limitations"]
    assert (tmp_path / "local_validation_manifest.json").is_file()
    assert not (tmp_path / "approved_manifest.json").exists()


def test_manifest_label_cannot_mark_generated_pack_distributable(tmp_path):
    source, manifest_path = _copy_source(tmp_path)
    source_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_manifest["classification"] = "SYNTHETIC_TEST_ONLY"
    manifest_path.write_text(json.dumps(source_manifest), encoding="utf-8")

    output = tmp_path / "local_validation_pack"
    result = anonymize_pack(source, manifest_path, output, as_of_utc=AS_OF)

    assert result["distribution_allowed"] is False
    assert result["classification"] == "LOCAL_VALIDATION_ONLY"
    assert not (tmp_path / "approved_manifest.json").exists()


def test_pack_cannot_be_written_to_arbitrary_tracked_project_path(tmp_path):
    source, manifest_path = _copy_source(tmp_path)
    output = Path(__file__).resolve().parents[2] / "backend" / "untracked_pack_name"

    with pytest.raises(AnonymizationError) as error:
        anonymize_pack(source, manifest_path, output, as_of_utc=AS_OF)

    assert error.value.code == "OUTPUT_PATH_NOT_PRIVATE"
    assert not output.exists()


def test_surrogates_do_not_reassign_codes_when_rows_reordered():
    first = Surrogates(secret=b"first-local-pack-secret")
    second = Surrogates(secret=b"first-local-pack-secret")
    first_codes = {value: first.get("employee_code", value, "SYN-EMP") for value in ("A", "B")}
    second_codes = {value: second.get("employee_code", value, "SYN-EMP") for value in ("B", "A")}

    assert first_codes == second_codes
    assert len(set(first_codes.values())) == 2
    assert first_codes["A"] != Surrogates(secret=b"different-local-secret").get("employee_code", "A", "SYN-EMP")


def test_stale_source_checksum_never_creates_pack(tmp_path):
    source, manifest_path = _copy_source(tmp_path)
    workbook_path = source / "01_cong_viec.xlsx"
    workbook = load_workbook(workbook_path)
    workbook["Dữ liệu"]["C2"] = "Changed title"
    workbook.save(workbook_path)
    workbook.close()
    output = tmp_path / "approved_pack"
    with pytest.raises(AnonymizationError) as error:
        anonymize_pack(source, manifest_path, output, as_of_utc=AS_OF)
    assert error.value.code == "SOURCE_CHECKSUM_MISMATCH"
    assert not output.exists() and not (tmp_path / "local_validation_manifest.json").exists()


def test_future_schedule_warning_is_retained_in_manifest(tmp_path):
    source, manifest_path = _copy_source(tmp_path)
    workbook_path = source / "06_ve_sinh.xlsx"
    workbook = load_workbook(workbook_path)
    data = workbook["Dữ liệu"]
    headers = [cell.value for cell in data[1]]
    data.cell(2, headers.index("scheduled_start") + 1).value = datetime(2027, 1, 1, 8)
    data.cell(2, headers.index("scheduled_end") + 1).value = datetime(2027, 1, 1, 10)
    workbook.save(workbook_path)
    workbook.close()
    _refresh_manifest_entry(manifest_path, workbook_path)
    output = tmp_path / "approved_pack"
    result = anonymize_pack(source, manifest_path, output, as_of_utc=AS_OF)
    assert result["source_preflight"]["status"] == "PASS"
    assert result["source_preflight"]["warning_codes"]["FUTURE_SCHEDULE"] == 2
    assert output.is_dir()


def test_future_terminal_event_never_creates_pack(tmp_path):
    source, manifest_path = _copy_source(tmp_path)
    workbook_path = source / "01_cong_viec.xlsx"
    workbook = load_workbook(workbook_path)
    data = workbook["Dữ liệu"]
    headers = [cell.value for cell in data[1]]
    data.cell(4, headers.index("completed_at") + 1).value = datetime(2099, 1, 1, 12)
    workbook.save(workbook_path)
    workbook.close()
    _refresh_manifest_entry(manifest_path, workbook_path)
    output = tmp_path / "approved_pack"
    with pytest.raises(AnonymizationError) as error:
        anonymize_pack(source, manifest_path, output, as_of_utc=AS_OF)
    assert error.value.code == "SOURCE_PREFLIGHT_BLOCKED"
    assert not output.exists() and not (tmp_path / "local_validation_manifest.json").exists()
